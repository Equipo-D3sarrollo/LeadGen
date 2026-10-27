"""reprocesar: marcas históricas de dato inválido y protección de diagnósticos (origen ia/reglas);
--rediagnosticar-antiguos; export sin recomendaciones ni puntuaciones para filas excluidas.

Solo bases/archivos temporales y clientes de IA simulados.
"""
import hashlib
import json
import os
import sqlite3
import subprocess
import sys
from types import SimpleNamespace

import openpyxl
import pandas as pd
import pytest

from leadgen import catalogo, database, diagnostico, export, google_maps, perfil, perfil_psicologico as pp, pipeline, reprocesar
from leadgen.scraper import EmpresaCruda

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CAT = catalogo.catalogo_por_defecto()
TEL_CORRUPTO = "5076507550745073507250715070"


def _fila(**extra):
    base = {"url_perfil": "https://dir.example/empresa/1", "nombre": "Empresa Uno", "telefono": "",
            "whatsapp": "3001234567", "facebook": "https://facebook.com/uno", "direccion": "Cra 1 #2-3, Pereira"}
    base.update(extra)
    return base


def _uno(fila, **kw):
    return reprocesar.reprocesar_filas([fila], CAT, **kw)[0]


class _ClienteIA:
    """Cliente de Anthropic simulado: devuelve `texto` o lanza error."""
    def __init__(self, texto="Diagnóstico IA simulado.", falla=False):
        self.texto, self.falla, self.messages = texto, falla, self

    def create(self, **_):
        if self.falla:
            raise RuntimeError("IA caída")
        return SimpleNamespace(content=[SimpleNamespace(text=self.texto)])


# --- 1. Marcas históricas de datos inválidos ---------------------------------------
def test_marca_historica_se_conserva_si_el_campo_sigue_vacio():
    con = _uno(_fila(datos_invalidos="telefono"), diagnosticar=False)
    sin = _uno(_fila(datos_invalidos=""), diagnosticar=False)
    assert con["datos_invalidos"] == "telefono" and sin["datos_invalidos"] == ""
    assert abs((sin["score_final"] - con["score_final"]) - 5) < 0.11   # penalización por dato inválido


def test_marca_historica_se_quita_si_el_dato_ahora_es_valido():
    assert _uno(_fila(telefono="3109876543", datos_invalidos="telefono"), diagnosticar=False)["datos_invalidos"] == ""


def test_solo_se_quita_la_marca_del_campo_corregido():
    r = _uno(_fila(correo="ventas@uno.co", datos_invalidos="telefono;correo"), diagnosticar=False)
    assert r["datos_invalidos"] == "telefono"


def test_dato_corrupto_actual_con_y_sin_historial_cuenta_una_vez():
    assert _uno(_fila(telefono=TEL_CORRUPTO), diagnosticar=False)["datos_invalidos"] == "telefono"
    assert _uno(_fila(telefono=TEL_CORRUPTO, datos_invalidos="telefono"), diagnosticar=False)["datos_invalidos"] == "telefono"


@pytest.mark.parametrize("valor", [float("nan"), None, ""])
def test_datos_invalidos_vacio_o_nan(valor):
    assert _uno(_fila(datos_invalidos=valor), diagnosticar=False)["datos_invalidos"] == ""


# --- 2. Origen del diagnóstico ----------------------------------------------------------
def _per_res():
    per = perfil.calcular_perfil(_fila(tiene_whatsapp=1, fuente="directorio"))
    return per, pp.analizar(per, CAT, 50)


def test_diagnosticar_con_origen():
    per, res = _per_res()
    e = {"nombre": "Empresa Uno"}
    assert diagnostico.diagnosticar_con_origen(per, res, e, usar_ia=True, cliente=_ClienteIA()) == ("Diagnóstico IA simulado.", "ia")
    assert diagnostico.diagnosticar_con_origen(per, res, e, usar_ia=True, cliente=_ClienteIA(falla=True))[1] == "reglas"
    assert diagnostico.diagnosticar_con_origen(per, res, e, usar_ia=True, cliente=_ClienteIA(texto="  "))[1] == "reglas"
    assert diagnostico.diagnosticar_con_origen(per, res, e, usar_ia=False, cliente=_ClienteIA())[1] == "reglas"
    # la API pública no cambia
    assert diagnostico.diagnosticar(per, res, e) == diagnostico.diagnostico_por_reglas(per, res, e)


@pytest.mark.parametrize("origen, previo, rediag, conserva", [
    ("ia", "TEXTO IA", False, True),
    ("ia", "TEXTO IA", True, True),        # la IA nunca se sobrescribe
    ("reglas", "VIEJO", False, False),
    (None, "ANTIGUO", False, True),        # origen desconocido: se conserva por defecto
    (None, "ANTIGUO", True, False),        # ... salvo --rediagnosticar-antiguos
    (None, "", False, False),              # sin texto: se genera
])
def test_reprocesar_protege_diagnosticos(origen, previo, rediag, conserva):
    r = _uno(_fila(diagnostico=previo, diagnostico_origen=origen), rediagnosticar_antiguos=rediag)
    if conserva:
        assert r["diagnostico"] == previo and r.get("diagnostico_origen") == origen
    else:
        assert r["diagnostico"] != previo and r["diagnostico"] and r["diagnostico_origen"] == "reglas"


def test_migracion_agrega_diagnostico_origen(tmp_path):
    db = str(tmp_path / "vieja.db")
    conn = sqlite3.connect(db)
    conn.execute(database.ESQUEMA)
    conn.execute("INSERT INTO empresas (url_perfil, nombre) VALUES ('u1', 'E')")
    conn.commit(); conn.close()
    with database.conectar(db) as c:
        assert "diagnostico_origen" in {r[1] for r in c.execute("PRAGMA table_info(empresas)")}
        assert c.execute("SELECT diagnostico_origen FROM empresas").fetchone()[0] is None


def _pipeline_maps(tmp_path, monkeypatch, nombre):
    crudas = [EmpresaCruda(nombre=nombre, url_perfil="https://www.google.com/maps/place/1",
                           direccion="Cra 1 #2-3, Pereira, Risaralda", telefono="3001234567")]
    monkeypatch.setattr(google_maps, "extraer_negocios_google_maps", lambda *a, **k: crudas)
    db = str(tmp_path / "ia.db")
    pipeline.ejecutar_pipeline("https://www.google.com/maps/search/x", usar_ia=True, db_path=db, verbose=False,
                               catalogo_path=str(tmp_path / "no.json"), consultar_dominio=False)
    with database.conectar(db) as conn:
        return dict(database.obtener_todas(conn)[0])


def test_pipeline_guarda_origen_ia(tmp_path, monkeypatch):
    import anthropic
    monkeypatch.setenv("ANTHROPIC_API_KEY", "clave-de-prueba")
    monkeypatch.setattr(anthropic, "Anthropic", lambda api_key: _ClienteIA())
    fila = _pipeline_maps(tmp_path, monkeypatch, "Clínica IA")
    assert fila["diagnostico"] == "Diagnóstico IA simulado." and fila["diagnostico_origen"] == "ia"


def test_pipeline_sin_clave_guarda_origen_reglas(tmp_path, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    fila = _pipeline_maps(tmp_path, monkeypatch, "Clínica Reglas")
    assert fila["diagnostico"] and fila["diagnostico_origen"] == "reglas"


# --- 3. CLI --rediagnosticar-antiguos ------------------------------------------------------
def test_cli_rediagnosticar_antiguos(tmp_path):
    db = str(tmp_path / "cli.db")
    with database.conectar(db) as conn:
        database.guardar_o_actualizar(conn, _fila(url_perfil="https://dir.example/empresa/a", nombre="Antigua",
                                                  diagnostico="ANTIGUO"))
        database.guardar_o_actualizar(conn, _fila(url_perfil="https://dir.example/empresa/b", nombre="Con IA",
                                                  diagnostico="TEXTO IA", diagnostico_origen="ia"))

    def correr(*extra):
        r = subprocess.run([sys.executable, "main.py", "--db", db, "--catalogo", str(tmp_path / "x.json"),
                            "reprocesar", *extra], capture_output=True, text=True, cwd=RAIZ)
        assert r.returncode == 0, r.stderr
        with database.conectar(db) as conn:
            return r.stdout, {f["nombre"]: dict(f) for f in database.obtener_todas(conn)}

    salida, filas = correr()
    assert filas["Antigua"]["diagnostico"] == "ANTIGUO" and filas["Antigua"]["diagnostico_origen"] is None
    assert filas["Con IA"]["diagnostico"] == "TEXTO IA"
    assert "1 diagn" in salida and "--rediagnosticar-antiguos" in salida

    salida, filas = correr("--rediagnosticar-antiguos")
    assert filas["Antigua"]["diagnostico"] != "ANTIGUO" and filas["Antigua"]["diagnostico_origen"] == "reglas"
    assert filas["Con IA"]["diagnostico"] == "TEXTO IA" and filas["Con IA"]["diagnostico_origen"] == "ia"


# --- 4. Export sin recomendaciones para excluidas -------------------------------------------
OCULTAS = export.COLUMNAS_COMERCIALES_OCULTAS_EXCLUIDAS
COMERCIAL = {"mejor_producto": "Diseño web", "compatibilidad_pct": 70.0, "fit_pct": 60.0, "intent_pct": 55.0,
             "compatibilidades": '{"diseno_web": {"compatibilidad_pct": 70.0}}', "pasa_filtro": 1,
             "score_final": 60.0, "lead_score": 75, "lead_clasificacion": "Cliente potencial bueno",
             "dim_madurez_digital": 30.0}


@pytest.fixture
def db_export(tmp_path):
    db = str(tmp_path / "export.db")
    per = perfil.calcular_perfil(_fila(fuente="directorio"))
    nueva = pp.a_columnas(per, pp.analizar(per, CAT, 75, es_empresa=False))
    with database.conectar(db) as conn:
        database.guardar_o_actualizar(conn, dict(COMERCIAL, url_perfil="v", nombre="Valida", es_empresa=1, fuera_de_mercado=0))
        database.guardar_o_actualizar(conn, dict(COMERCIAL, url_perfil="l", nombre="Antigua Excluida", es_empresa=0, fuera_de_mercado=0))
        database.guardar_o_actualizar(conn, dict(COMERCIAL, url_perfil="f", nombre="Fuera", es_empresa=1, fuera_de_mercado=1))
        database.guardar_o_actualizar(conn, dict(nueva, url_perfil="n", nombre="Nueva Excluida", es_empresa=0, fuera_de_mercado=0))
        database.registrar_contacto(conn, "v", "diseno_web", "interesado")
    return db


def _leer_exports(db, tmp_path):
    with database.conectar(db) as conn:
        filas, contactos = database.obtener_todas(conn), database.obtener_contactos(conn)
    xlsx = export.exportar_excel(filas, str(tmp_path / "e.xlsx"), contactos=contactos)
    csv = export.exportar_csv(filas, str(tmp_path / "e.csv"))
    js = export.exportar_json(filas, str(tmp_path / "e.json"))
    hojas = pd.read_excel(xlsx, sheet_name=None)
    assert set(hojas) == {"Leads", "Contactos"}
    return {
        "excel": {r["nombre"]: r for r in hojas["Leads"].to_dict("records")},
        "csv": {r["nombre"]: r for r in pd.read_csv(csv).to_dict("records")},
        "json": {r["nombre"]: r for r in json.load(open(js, encoding="utf-8"))},
    }


def test_export_oculta_recomendaciones_de_excluidas(db_export, tmp_path):
    for formato, filas in _leer_exports(db_export, tmp_path).items():
        valida = filas["Valida"]
        assert valida["mejor_producto"] == "Diseño web" and valida["score_final"] == 60.0, formato
        for nombre in ("Antigua Excluida", "Fuera", "Nueva Excluida"):
            fila = filas[nombre]
            assert all(pd.isna(fila[c]) for c in OCULTAS), (formato, nombre, {c: fila[c] for c in OCULTAS})
        # el perfil descriptivo y las banderas de exclusión se mantienen
        assert filas["Antigua Excluida"]["dim_madurez_digital"] == 30.0 and filas["Antigua Excluida"]["es_empresa"] == 0
        assert filas["Fuera"]["fuera_de_mercado"] == 1


def test_export_no_altera_la_base(db_export, tmp_path):
    with open(db_export, "rb") as f:
        antes = hashlib.sha256(f.read()).hexdigest()
    _leer_exports(db_export, tmp_path)
    with open(db_export, "rb") as f:
        assert hashlib.sha256(f.read()).hexdigest() == antes
    with database.conectar(db_export) as conn:
        guardada = conn.execute("SELECT mejor_producto, score_final FROM empresas WHERE url_perfil='l'").fetchone()
    assert guardada["mejor_producto"] == "Diseño web" and guardada["score_final"] == 60.0


def test_export_tolera_filas_sin_columnas_de_exclusion(tmp_path):
    ruta = export.exportar_csv([{"nombre": "Sin banderas", "score_final": 50.0, "fuera_de_mercado": float("nan")}],
                               str(tmp_path / "s.csv"))
    assert pd.read_csv(ruta).loc[0, "score_final"] == 50.0


def test_export_protege_formulas_y_conserva_texto_en_excel(tmp_path):
    valores = ["=1+1", " +SUM(A1)", "\t-10+2", "\t@cmd", "+18187534867", "texto normal", 12.5]
    filas = [{"nombre": valor, "telefono": valor} for valor in valores]
    ruta = export.exportar_excel(filas, str(tmp_path / "seguro.xlsx"))

    libro = openpyxl.load_workbook(ruta, data_only=False)
    hoja = libro["Leads"]
    nombres = [hoja.cell(fila, 1).value for fila in range(2, len(valores) + 2)]
    telefonos = [hoja.cell(fila, 2).value for fila in range(2, len(valores) + 2)]
    assert nombres == valores
    assert telefonos == valores
    assert all(hoja.cell(fila, 1).data_type == "s" for fila in range(2, 6))


def test_export_neutraliza_formulas_en_csv_sin_alterar_telefonos_ni_numeros(tmp_path):
    valores = ["=1+1", " +SUM(A1)", "\t-10+2", "\x00@cmd", "+18187534867", "texto normal", 12.5, -10]
    ruta = export.exportar_csv(
        [{"nombre": valor, "telefono": valor, "numero": valor} for valor in valores],
        str(tmp_path / "seguro.csv"),
    )
    exportado = pd.read_csv(ruta, dtype={"nombre": str, "telefono": str, "numero": str})

    assert exportado["nombre"].tolist() == ["'=1+1", "' +SUM(A1)", "'\t-10+2", "'",
                                            "+18187534867", "texto normal", "12.5", "-10"]
    assert exportado["telefono"].tolist() == exportado["nombre"].tolist()
    assert exportado["numero"].tolist() == exportado["nombre"].tolist()
    assert "'\x00@cmd" in open(ruta, encoding="utf-8-sig").read()
